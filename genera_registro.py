#!/usr/bin/env python3
"""
genera_registro.py
Genera un file Word con il registro degli articoli (tabella progressivo/data/titolo).

Utilizzo:
    # Da vault Obsidian:
    python3 genera_registro.py --vault /percorso/vault --output registro.docx

    # Da esportazione WordPress:
    python3 genera_registro.py --xml wordpress-export.xml --output registro.docx

    # Con logo personalizzato:
    python3 genera_registro.py --vault /percorso/vault --logo logo.png --output registro.docx

Requisiti:
    pip install python-docx
"""

import sys
import argparse
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("ERRORE: python-docx non è installato.", file=sys.stderr)
    print("       Installa con: pip install python-docx", file=sys.stderr)
    sys.exit(1)


# ─── Colori ───────────────────────────────────────────────────────────────────

COLOR_HEADER = RGBColor(0xA9, 0xD1, 0x8E)   # verde intestazione
COLOR_ODD    = RGBColor(0xE2, 0xEF, 0xDA)   # verde righe dispari
COLOR_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)


# ─── Helpers XML ──────────────────────────────────────────────────────────────

def _set_cell_bg(cell, rgb: RGBColor):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}")
    tcPr.append(shd)


# ─── Costruzione documento ────────────────────────────────────────────────────

HEADERS = [
    "Numero\nProgressivo\ndegli articoli\npresentati",
    "Data",
    "Titolo articolo",
    "Pagina",
    "Descrizione\nArt. (*)",
]
COL_WIDTHS = [Cm(3.2), Cm(2.6), Cm(8.2), Cm(2.0), Cm(2.6)]
COL_ALIGNS = [
    WD_ALIGN_PARAGRAPH.CENTER,
    WD_ALIGN_PARAGRAPH.CENTER,
    WD_ALIGN_PARAGRAPH.LEFT,
    WD_ALIGN_PARAGRAPH.CENTER,
    WD_ALIGN_PARAGRAPH.CENTER,
]


def build_document(articles: list[dict], logo_path: Path | None = None) -> Document:
    doc = Document()

    # Margini pagina
    section = doc.sections[0]
    section.top_margin    = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin   = Cm(2.0)
    section.right_margin  = Cm(2.0)

    # ── Header: logo o testo ──────────────────────────────────────────────────
    if logo_path and logo_path.exists():
        p    = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run  = p.add_run()
        run.add_picture(str(logo_path), width=Cm(6))
    else:
        p    = doc.add_paragraph("WINE CHANNEL")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r    = p.runs[0]
        r.bold           = True
        r.font.size      = Pt(20)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    sub   = doc.add_paragraph("Direttore Responsabile: Giancarlo Febbo")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs    = sub.runs[0]
    rs.bold       = True
    rs.font.size  = Pt(10)

    doc.add_paragraph()  # spazio prima della tabella

    # ── Tabella ──────────────────────────────────────────────────────────────
    table = doc.add_table(rows=1 + len(articles), cols=5)
    table.style = "Table Grid"

    # Intestazione
    hdr_row = table.rows[0]
    for i, (text, width) in enumerate(zip(HEADERS, COL_WIDTHS)):
        cell = hdr_row.cells[i]
        cell.width = width
        _set_cell_bg(cell, COLOR_HEADER)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold       = True
        r.font.size  = Pt(9)

    # Righe dati
    for row_idx, article in enumerate(articles, start=1):
        bg = COLOR_ODD if row_idx % 2 == 1 else COLOR_WHITE

        pub_date = article.get("pub_date")
        date_str = (
            pub_date.strftime("%d/%m/%Y")
            if pub_date and pub_date != datetime.min
            else ""
        )
        title = article.get("title", "").upper()

        values = [str(row_idx), date_str, title, "/", "F"]

        row = table.rows[row_idx]
        for i, (val, align, width) in enumerate(zip(values, COL_ALIGNS, COL_WIDTHS)):
            cell = row.cells[i]
            cell.width = width
            _set_cell_bg(cell, bg)
            p = cell.paragraphs[0]
            p.alignment = align
            r = p.add_run(val)
            r.font.size = Pt(9)

    return doc


# ─── Sorgenti dati ────────────────────────────────────────────────────────────

def load_from_vault(vault: Path) -> list[dict]:
    try:
        from obsidian_to_wordpress import collect_articles
    except ImportError:
        print("ERRORE: obsidian_to_wordpress.py non trovato nella stessa cartella.", file=sys.stderr)
        sys.exit(1)
    articles = collect_articles(vault)
    return sorted(articles, key=lambda a: a["pub_date"])


def load_from_xml(xml_path: Path) -> list[dict]:
    try:
        from wordpress_to_pdf import parse_articles_from_wxr
    except ImportError:
        print("ERRORE: wordpress_to_pdf.py non trovato nella stessa cartella.", file=sys.stderr)
        sys.exit(1)
    articles, _ = parse_articles_from_wxr(xml_path)
    return sorted(articles, key=lambda a: a["pub_date"])


# ─── CLI ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Genera il registro Word degli articoli (tabella progressivo/data/titolo)",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Esempi:
  python3 genera_registro.py --vault ~/Obsidian/Blog --output registro.docx
  python3 genera_registro.py --xml wordpress-export.xml --output registro.docx
  python3 genera_registro.py --xml wordpress-export.xml --logo winechannel.png --output registro.docx

Requisiti:
  pip install python-docx
""",
    )

    src = parser.add_mutually_exclusive_group(required=True)
    src.add_argument("--vault", "-v", help="Percorso vault Obsidian")
    src.add_argument("--xml",   "-x", help="File XML di esportazione WordPress")

    parser.add_argument("--output", "-o", default="registro.docx",
                        help="File Word di output (default: registro.docx)")
    parser.add_argument("--logo", "-l", default=None,
                        help="Immagine logo da inserire nell'intestazione (PNG/JPG)")

    args = parser.parse_args()

    logo_path = Path(args.logo).resolve() if args.logo else None

    if args.vault:
        vault = Path(args.vault).expanduser().resolve()
        if not vault.is_dir():
            print(f"ERRORE: '{vault}' non è una cartella valida.", file=sys.stderr)
            sys.exit(1)
        print(f"Lettura vault: {vault}")
        articles = load_from_vault(vault)
    else:
        xml_path = Path(args.xml).expanduser().resolve()
        if not xml_path.is_file():
            print(f"ERRORE: '{xml_path}' non trovato.", file=sys.stderr)
            sys.exit(1)
        print(f"Lettura XML: {xml_path}")
        articles = load_from_xml(xml_path)

    if not articles:
        print("Nessun articolo trovato.", file=sys.stderr)
        sys.exit(1)

    print(f"Articoli trovati: {len(articles)}")
    print("Generazione documento Word...", end=" ", flush=True)

    doc      = build_document(articles, logo_path=logo_path)
    out_path = Path(args.output).expanduser().resolve()
    doc.save(str(out_path))

    print(f"✓\nFile salvato: {out_path}")


if __name__ == "__main__":
    main()
